"""docs/rapor/ altındaki bölümleri birleştirip Word ve PDF rapor üretir.

    python -m scripts.rapor_olustur            # docx + pdf
    python -m scripts.rapor_olustur --no-pdf   # yalnızca docx
    python -m scripts.rapor_olustur --belge ../docs/kullanici_testi.md   # tek belge: yanına .docx ve .pdf

Çıktı: docs/rapor/cikti/Bitirme_Raporu.docx ve Bitirme_Raporu.pdf.
İçindekiler ve PDF için Microsoft Word gerekir (pywin32 üzerinden). Word yoksa docx
yine üretilir; içindekiler Word'de açıldığında güncellenir.

Bağımlılıklar: requirements-rapor.txt (python-docx, pywin32; görseller için matplotlib).
Markdown'ın yalnızca rapor dosyalarında kullandığımız alt kümesi desteklenir: başlıklar,
paragraflar, çitli kod blokları (```), madde ve numaralı listeler (iki boşluk girintili alt maddeler), tablolar,
görseller, alıntı blokları, kalın/italik/kod ve [DOLDUR: ...] işaretleri (sarı vurgulanır).
Yazdırılacak belgelerde `<div style="page-break-after: always"></div>` satırı sayfa sonu olur, `---` atlanır.
"""

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

RAPOR = Path(__file__).resolve().parents[2] / "docs" / "rapor"
CIKTI = RAPOR / "cikti"
BOLUMLER = [
    "00_ozet.md",
    "01_hazirlik_ve_tasarim.md",
    "02_veri_toplama.md",
    "03_backend_ve_metrikler.md",
    "04_frontend.md",
    "05_risk_profili_ve_oneri.md",
    "06_ml_tahmin.md",
    "07_test_ve_demo.md",
    "08_sonuc.md",
    "kaynakca.md",
]
KAPAK = [
    ("Fon Analiz Platformu", 26, True),
    ("Bitirme Projesi Raporu", 16, False),
    ("", 12, False),
    ("Mehmet Korkut", 16, True),
    ("Bilgisayar Mühendisliği", 14, False),
    ("[DOLDUR: üniversite]", 14, False),
    ("", 12, False),
    ("Danışman: [DOLDUR: danışman]", 13, False),
    ("[DOLDUR: tarih]", 13, False),
]
FONT = "Times New Roman"
KOD_FONT = "Consolas"
MAX_GENISLIK = Cm(15.5)
MAX_YUKSEKLIK = Cm(17)

INLINE = re.compile(
    r"(\*\*.+?\*\*"                    # kalın
    r"|`[^`]+`"                         # kod
    r"|\[DOLDUR[^\]]*\]"                # doldurulacak yer
    r"|\[[^\]]+\]\([^)]+\)"             # bağlantı
    r"|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))"  # italik
)


# --- satır içi biçimler ----------------------------------------------------------------

def satir_ici(p, text, bold=False, italic=False, size=None):
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            _run(p, text[pos:m.start()], bold, italic, size)
        tok = m.group(0)
        if tok.startswith("**"):
            satir_ici(p, tok[2:-2], True, italic, size)
        elif tok.startswith("`"):
            r = _run(p, tok[1:-1], bold, italic, size)
            r.font.name = KOD_FONT
            r.font.size = Pt((size or 12) - 1.5)
        elif tok.startswith("[DOLDUR"):
            r = _run(p, tok, True, italic, size)
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif tok.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok).groups()
            _run(p, label if label == url else f"{label} ({url})", bold, italic, size)
        else:
            satir_ici(p, tok[1:-1], bold, True, size)
        pos = m.end()
    if pos < len(text):
        _run(p, text[pos:], bold, italic, size)


def _run(p, text, bold, italic, size):
    r = p.add_run(text)
    r.bold = bold or None
    r.italic = italic or None
    if size:
        r.font.size = Pt(size)
    return r


# --- belge kurulumu ----------------------------------------------------------------------

def stiller(doc):
    st = doc.styles
    normal = st["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    _dogu_asya_font(normal)
    _dil(normal, "tr-TR")
    pf = normal.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.3
    for name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)]:
        s = st[name]
        s.font.name = FONT
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.italic = False
        s.font.color.rgb = RGBColor(0, 0, 0)
        _dogu_asya_font(s)
        s.paragraph_format.space_before = Pt(18 if name == "Heading 1" else 12)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True
    for name in ["List Bullet", "List Bullet 2", "List Number"]:
        st[name].paragraph_format.space_after = Pt(3)


def _dogu_asya_font(style):
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(a), None)


def _dil(style, lang):
    rpr = style.element.get_or_add_rPr()
    el = OxmlElement("w:lang")
    el.set(qn("w:val"), lang)
    rpr.append(el)


def alan(p, instr, yer_tutucu=""):
    """Paragrafa bir Word alanı (TOC, PAGE) ekler."""
    r = p.add_run()
    for tag, attr in [("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "separate")]:
        el = OxmlElement(tag)
        if attr:
            el.set(qn("w:fldCharType"), attr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        r._r.append(el)
    if yer_tutucu:
        p.add_run(yer_tutucu)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    p.add_run()._r.append(end)


def sayfa_duzeni(doc):
    s = doc.sections[0]
    s.page_height, s.page_width = Cm(29.7), Cm(21)
    s.left_margin = s.right_margin = Cm(2.5)
    s.top_margin = s.bottom_margin = Cm(2.5)
    s.different_first_page_header_footer = True
    p = s.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    alan(p, "PAGE")


def kapak(doc):
    for _ in range(5):
        doc.add_paragraph()
    for text, size, bold in KAPAK:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        satir_ici(p, text, bold=bold, size=size)
    sayfa_sonu(doc)


def icindekiler(doc):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    satir_ici(p, "İçindekiler", bold=True, size=16)
    alan(doc.add_paragraph(), r'TOC \o "1-2" \h \z \u',
         "İçindekiler tablosu Word'de alan güncellenince oluşur (F9).")


def sayfa_sonu(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# --- markdown bloklar --------------------------------------------------------------------

PAGE_BREAK = '<div style="page-break-after: always"></div>'


def bloklar(lines):
    """Satırları (tür, veri) bloklarına ayırır."""
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip() or line.strip() == "---":
            i += 1
            continue
        if line.strip() == PAGE_BREAK:
            yield "sayfa_sonu", None
            i += 1
        elif line.strip().startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            i += 1  # kapanış çiti
            yield "kod", code
        elif m := re.match(r"^(#{1,3}) (.+)$", line):
            yield "baslik", (len(m.group(1)), m.group(2))
            i += 1
        elif m := re.match(r"^!\[(.*)\]\((.+)\)$", line.strip()):
            yield "gorsel", (m.group(1), m.group(2))
            i += 1
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i].rstrip())
                i += 1
            yield "tablo", rows
        elif line.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].startswith(">"):
                buf.append(lines[i].lstrip(">").strip())
                i += 1
            yield "alinti", " ".join(b for b in buf if b)
        elif re.match(r"^\s*(- |\d+\. )", line):
            items = []
            while i < len(lines) and lines[i].strip():
                ln = lines[i].rstrip()
                if m := re.match(r"^(\s*)(- |(\d+)\. )(.*)$", ln):
                    level = len(m.group(1)) // 2
                    items.append([level, m.group(3), m.group(4)])
                elif items:  # önceki maddenin devamı
                    items[-1][2] += " " + ln.strip()
                i += 1
            yield "liste", items
        else:
            buf = []
            while i < len(lines) and lines[i].strip() and lines[i].strip() not in (PAGE_BREAK, "---") and not lines[i].strip().startswith("```") and not re.match(
                r"^(#{1,3} |!\[|\||>|\s*- |\s*\d+\. )", lines[i]
            ):
                buf.append(lines[i].strip())
                i += 1
            yield "paragraf", " ".join(buf)


def hucreler(row):
    return [c.strip() for c in row.strip().strip("|").split("|")]


def tablo(doc, rows):
    rows = [r for r in rows if not re.match(r"^\|[\s:|-]+\|$", r)]
    data = [hucreler(r) for r in rows]
    ncol = max(len(r) for r in data)
    t = doc.add_table(rows=len(data), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    bos_baslik = all(not c for c in data[0])
    for ri, row in enumerate(data):
        for ci in range(ncol):
            cell = t.cell(ri, ci)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            text = row[ci] if ci < len(row) else ""
            satir_ici(p, text, bold=(ri == 0 and not bos_baslik), size=10)
            if ri == 0 and not bos_baslik:
                _golge(cell, "E8EEF7")
    if bos_baslik:  # "| | |" başlıklı anahtar-değer tabloları
        t._tbl.remove(t.rows[0]._tr)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def _golge(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def gorsel(doc, alt, path, base):
    img = (base / path).resolve()
    if not img.exists():
        p = doc.add_paragraph()
        satir_ici(p, f"[DOLDUR: görsel bulunamadı: {path}]")
        return
    w, h = Image.open(img).size
    width = MAX_GENISLIK
    if width * h / w > MAX_YUKSEKLIK:
        width = int(MAX_YUKSEKLIK * w / h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(img), width=width)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    m = re.match(r"^(Şekil [\d.]+)\s*(.*)$", alt)
    if m:
        _run(cap, m.group(1) + " ", True, False, 10)
        satir_ici(cap, m.group(2), size=10)
    else:
        satir_ici(cap, alt, size=10)


def liste(doc, items):
    for level, num, text in items:
        if num:
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.9 + 0.8 * level)
            pf.first_line_indent = Cm(-0.6)
            pf.space_after = Pt(3)
            _run(p, f"{num}.\t", False, False, None)
            pf.tab_stops.add_tab_stop(pf.left_indent)
        else:
            p = doc.add_paragraph(style="List Bullet 2" if level else "List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        satir_ici(p, text)


def bolum(doc, md, ilk):
    base = md.parent
    for kind, data in bloklar(md.read_text(encoding="utf-8").splitlines()):
        if kind == "baslik":
            level, text = data
            h = doc.add_heading(text, level=level)
            if level == 1 and not ilk:
                h.paragraph_format.page_break_before = True
            ilk = False
        elif kind == "paragraf":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if re.match(r"^\*\*Tablo [\d.]+\*\*", data):
                p.paragraph_format.keep_with_next = True
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            satir_ici(p, data)
        elif kind == "alinti":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.right_indent = Cm(1)
            satir_ici(p, data, italic=True, size=11)
        elif kind == "liste":
            liste(doc, data)
        elif kind == "tablo":
            tablo(doc, data)
        elif kind == "gorsel":
            gorsel(doc, *data, base)
        elif kind == "sayfa_sonu":
            sayfa_sonu(doc)
        elif kind == "kod":
            for ln in data:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(ln or " ")
                r.font.name = KOD_FONT
                r.font.size = Pt(9)


# --- Word: içindekiler ve PDF ------------------------------------------------------------

def word_ile_guncelle(docx, pdf):
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        print("pywin32 yok: içindekiler güncellenmedi, PDF üretilmedi.")
        return False
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        d = word.Documents.Open(str(docx))
        for style_id in (-20, -21):  # wdStyleTOC1, wdStyleTOC2: sıkı satır aralığı
            pf = d.Styles(style_id).ParagraphFormat
            pf.SpaceBefore, pf.SpaceAfter, pf.LineSpacingRule = 0, 2, 0
        d.Styles(-20).Font.Bold = True
        for i in range(1, d.TablesOfContents.Count + 1):
            d.TablesOfContents(i).Update()
        d.Fields.Update()
        d.Save()
        if pdf:
            d.ExportAsFixedFormat(str(pdf), 17)  # wdExportFormatPDF
        d.Close(False)
        return True
    except Exception as e:  # Word yok ya da açılamadı
        print(f"Word ile güncelleme başarısız: {e}")
        return False
    finally:
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass  # Word bazen kapanırken RPC hatası verir; belge zaten kaydedildi


def guncelleme_iste(docx):
    """Word yoksa, belge açılınca alanları güncellemesini iste."""
    doc = Document(str(docx))
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    doc.settings.element.append(el)
    doc.save(str(docx))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--no-pdf", action="store_true", help="PDF üretme")
    ap.add_argument("--belge", type=Path, help="Tek bir Markdown belgesini (kapak ve içindekiler olmadan) yanına docx/pdf olarak çevir")
    args = ap.parse_args(argv)

    if args.belge:
        md = args.belge.resolve()
        doc = Document()
        stiller(doc)
        sayfa_duzeni(doc)
        bolum(doc, md, ilk=True)
        docx, pdf = md.with_suffix(".docx"), md.with_suffix(".pdf")
        doc.save(str(docx))
        print(f"docx: {docx}")
        if not args.no_pdf and word_ile_guncelle(docx, pdf):
            print(f"pdf:  {pdf}")
        return

    CIKTI.mkdir(parents=True, exist_ok=True)
    docx = CIKTI / "Bitirme_Raporu.docx"
    pdf = CIKTI / "Bitirme_Raporu.pdf"

    doc = Document()
    stiller(doc)
    sayfa_duzeni(doc)
    kapak(doc)
    for i, name in enumerate(BOLUMLER):
        if name == "01_hazirlik_ve_tasarim.md":
            icindekiler(doc)
        bolum(doc, RAPOR / name, ilk=i == 0)
    try:
        doc.save(str(docx))
    except PermissionError:
        sys.exit(f"{docx} yazılamadı; Word'de açıksa kapatın.")
    print(f"docx: {docx}")

    pdf_yaz = not args.no_pdf
    if pdf_yaz and pdf.exists():
        try:
            pdf.open("ab").close()
        except PermissionError:
            print(f"{pdf} bir görüntüleyicide açık; PDF üretilmedi, kapatıp yeniden çalıştırın.")
            pdf_yaz = False

    if word_ile_guncelle(docx, pdf if pdf_yaz else None):
        if pdf_yaz:
            print(f"pdf:  {pdf}")
    else:
        guncelleme_iste(docx)


if __name__ == "__main__":
    main()
